"""Read-only provider access. Tokens never leave the local backend."""
import base64
import hashlib
import io
import json
import os
import re
import secrets
import time
import zipfile
from urllib.parse import urlencode, urlparse, quote

import httpx
from cryptography.fernet import Fernet, InvalidToken
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from fastapi import HTTPException

PROVIDERS = {
    'google': {'authorize': 'https://accounts.google.com/o/oauth2/v2/auth', 'token': 'https://oauth2.googleapis.com/token', 'scope': 'https://www.googleapis.com/auth/drive.readonly'},
    'microsoft': {'authorize': 'https://login.microsoftonline.com/common/oauth2/v2.0/authorize', 'token': 'https://login.microsoftonline.com/common/oauth2/v2.0/token', 'scope': 'https://graph.microsoft.com/Files.Read offline_access'},
}
LIMIT = 20 * 1024 * 1024

class Accounts:
    def __init__(self, security):
        self.security = security

    def config(self, provider):
        if provider not in PROVIDERS:
            raise HTTPException(400, 'Unknown provider')
        prefix = 'STARVIEW_' + provider.upper()
        client = os.environ.get(prefix + '_CLIENT_ID', '')
        secret = os.environ.get(prefix + '_CLIENT_SECRET', '')
        if not client or not secret:
            raise HTTPException(503, f'{provider.title()} connection needs OAuth app credentials. See the setup guide.')
        callback = os.environ.get('STARVIEW_API_ORIGIN', 'http://127.0.0.1:8000').rstrip('/') + '/api/accounts/' + provider + '/callback'
        return {**PROVIDERS[provider], 'client_id': client, 'client_secret': secret, 'redirect_uri': callback}

    def cipher(self):
        path = self.security.directory / 'token.key'
        with self.security.lock:
            if not path.exists():
                fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
                with os.fdopen(fd, 'wb') as file:
                    file.write(Fernet.generate_key())
            return Fernet(path.read_bytes())

    def statuses(self):
        with self.security.db() as db:
            connected = {row['provider'] for row in db.execute('SELECT provider FROM accounts')}
        return [{'provider': name, 'connected': name in connected, 'configured': bool(os.environ.get('STARVIEW_' + name.upper() + '_CLIENT_ID') and os.environ.get('STARVIEW_' + name.upper() + '_CLIENT_SECRET'))} for name in PROVIDERS]

    def start(self, provider, session):
        config = self.config(provider)
        state, verifier = secrets.token_urlsafe(32), secrets.token_urlsafe(64)
        with self.security.db() as db:
            db.execute('DELETE FROM oauth_states WHERE expires<?', (time.time(),))
            db.execute('INSERT INTO oauth_states VALUES (?,?,?,?,?)', (self.security.digest(state), session, provider, verifier, time.time() + 600))
        params = {key: config[key] for key in ('client_id', 'redirect_uri', 'scope')}
        params.update(response_type='code', state=state, code_challenge=base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).decode().rstrip('='), code_challenge_method='S256', prompt='consent' if provider == 'google' else 'select_account')
        if provider == 'google':
            params['access_type'] = 'offline'
        return config['authorize'] + '?' + urlencode(params)

    def finish(self, provider, state, code, session):
        config = self.config(provider)
        with self.security.lock, self.security.db() as db:
            row = db.execute('SELECT * FROM oauth_states WHERE digest=?', (self.security.digest(state),)).fetchone()
            if not row or row['provider'] != provider or row['session'] != session or row['expires'] < time.time():
                raise HTTPException(400, 'Account connection expired or could not be verified')
            db.execute('DELETE FROM oauth_states WHERE digest=?', (self.security.digest(state),))
        payload = {key: config[key] for key in ('client_id', 'client_secret', 'redirect_uri')}
        payload.update(grant_type='authorization_code', code=code, code_verifier=row['verifier'])
        tokens = self.exchange(config['token'], payload)
        self.save(provider, tokens)

    def exchange(self, url, payload):
        try:
            response = httpx.post(url, data=payload, timeout=20, follow_redirects=False)
            if response.status_code != 200:
                raise HTTPException(502, 'Account authorization failed. Reconnect and try again.')
            tokens = response.json()
            if not isinstance(tokens.get('access_token'), str):
                raise HTTPException(502, 'Provider did not return an access token')
            tokens['expires_at'] = time.time() + int(tokens.get('expires_in', 3600))
            return tokens
        except (httpx.HTTPError, ValueError):
            raise HTTPException(502, 'Could not reach the account provider')

    def save(self, provider, tokens):
        encrypted = self.cipher().encrypt(json.dumps(tokens).encode())
        with self.security.db() as db:
            db.execute('INSERT OR REPLACE INTO accounts VALUES (?,?)', (provider, encrypted))

    def token(self, provider):
        with self.security.lock:
            with self.security.db() as db:
                row = db.execute('SELECT tokens FROM accounts WHERE provider=?', (provider,)).fetchone()
            if not row:
                raise HTTPException(409, 'Connect this account to view its resources')
            try:
                tokens = json.loads(self.cipher().decrypt(row['tokens']))
            except (InvalidToken, ValueError):
                raise HTTPException(409, 'Account credentials could not be read. Disconnect and reconnect this account.')
            if tokens['expires_at'] < time.time() + 60:
                if not tokens.get('refresh_token'):
                    raise HTTPException(409, 'Your account connection expired. Reconnect it.')
                config = self.config(provider)
                payload = {key: config[key] for key in ('client_id', 'client_secret')}
                payload.update(grant_type='refresh_token', refresh_token=tokens['refresh_token'])
                if provider == 'microsoft':
                    payload['scope'] = config['scope']
                new = self.exchange(config['token'], payload)
                new.setdefault('refresh_token', tokens['refresh_token'])
                self.save(provider, new)
                tokens = new
            return tokens['access_token']

    def disconnect(self, provider):
        if provider not in PROVIDERS:
            raise HTTPException(400, 'Unknown provider')
        with self.security.lock, self.security.db() as db:
            db.execute('DELETE FROM accounts WHERE provider=?', (provider,))
            db.execute('DELETE FROM oauth_states WHERE provider=?', (provider,))

    @staticmethod
    def resource_id(value):
        if not re.fullmatch(r'[A-Za-z0-9_.!-]{1,200}', value):
            raise HTTPException(400, 'Invalid resource identifier')
        return quote(value, safe='')

    def fetch(self, provider, path, params=None, binary=False):
        base = 'https://www.googleapis.com/drive/v3' if provider == 'google' else 'https://graph.microsoft.com/v1.0'
        headers = {'Authorization': 'Bearer ' + self.token(provider)}
        try:
            with httpx.Client(timeout=30, follow_redirects=False) as client:
                with client.stream('GET', base + path, params=params, headers=headers) as response:
                    if response.status_code in {301, 302, 303, 307, 308} and binary:
                        location = response.headers.get('location', '')
                        parsed = urlparse(location)
                        host = parsed.hostname or ''
                        if parsed.scheme != 'https' or parsed.username or not any(host == suffix or host.endswith('.' + suffix) for suffix in ('1drv.com', 'onedrive.com', 'sharepoint.com', 'sharepoint.cn', 'files.1drv.com')):
                            raise HTTPException(502, 'Provider returned an unsupported download destination')
                        # A provider-issued signed URL, fetched without forwarding the bearer token.
                        with client.stream('GET', location) as download:
                            return self.body(download, True)
                    return self.body(response, binary)
        except httpx.HTTPError:
            raise HTTPException(502, 'Could not load this resource from the provider')

    @staticmethod
    def body(response, binary):
        if response.status_code != 200:
            raise HTTPException(502, 'Resource is unavailable, permission was denied, or the account needs reconnecting')
        data = bytearray()
        for chunk in response.iter_bytes():
            data.extend(chunk)
            if len(data) > LIMIT:
                raise HTTPException(413, 'Preview supports files up to 20 MB')
        if binary:
            return bytes(data)
        try:
            return json.loads(data)
        except ValueError:
            raise HTTPException(502, 'Provider returned invalid data')

    def resource(self, provider, resource_id):
        if provider not in PROVIDERS:
            raise HTTPException(400, 'Unknown provider')
        key = self.resource_id(resource_id)
        if provider == 'google':
            item = self.fetch(provider, '/files/' + key, {'fields': 'id,name,mimeType'})
            mime = item.get('mimeType', '')
            kind = 'sheet' if mime in {'application/vnd.google-apps.spreadsheet', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'} else 'document' if mime in {'application/vnd.google-apps.document', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/pdf'} else None
        else:
            item = self.fetch(provider, '/me/drive/items/' + key, {'$select': 'id,name,file'})
            name = item.get('name', '').lower()
            kind = 'sheet' if name.endswith('.xlsx') else 'document' if name.endswith(('.docx', '.pdf')) else None
            mime = item.get('file', {}).get('mimeType', '')
        if not kind:
            raise HTTPException(400, 'Choose a Google document, spreadsheet, Word document, Excel workbook, or PDF')
        return {'provider': provider, 'id': item['id'], 'name': item['name'], 'kind': kind, 'mime': mime}

    def browse(self, provider, folder='', cursor=''):
        if provider not in PROVIDERS:
            raise HTTPException(400, 'Unknown provider')
        if len(cursor) > 4096:
            raise HTTPException(400, 'Invalid page cursor')
        if provider == 'google':
            folder_id = self.resource_id(folder or 'root')
            params = {'q': f"'{folder_id}' in parents and trashed=false", 'fields': 'nextPageToken,files(id,name,mimeType)', 'pageSize': 100, 'orderBy': 'folder,name'}
            if cursor:
                params['pageToken'] = cursor
            result = self.fetch(provider, '/files', params)
            raw = result.get('files', [])
            next_cursor = result.get('nextPageToken', '')
        else:
            path = '/me/drive/items/' + self.resource_id(folder) + '/children' if folder else '/me/drive/root/children'
            params = {'$select': 'id,name,file,folder', '$top': 100}
            if cursor:
                params['$skiptoken'] = cursor
            result = self.fetch(provider, path, params)
            raw = result.get('value', [])
            # Never follow a provider URL supplied by a client; keep only its pagination token.
            from urllib.parse import parse_qs
            next_cursor = parse_qs(urlparse(result.get('@odata.nextLink', '')).query).get('$skiptoken', [''])[0]
        items = []
        for item in raw:
            folder_item = item.get('mimeType') == 'application/vnd.google-apps.folder' or 'folder' in item
            name = item.get('name', '')
            mime = item.get('mimeType', '')
            if folder_item or mime in {'application/vnd.google-apps.document', 'application/vnd.google-apps.spreadsheet'} or name.lower().endswith(('.docx', '.xlsx', '.pdf')):
                items.append({'id': item['id'], 'name': name, 'folder': folder_item})
        return {'items': items, 'cursor': next_cursor}

    def preview(self, resource):
        # Recheck metadata: no caller-provided URL or MIME controls an outbound request.
        resource = self.resource(resource['provider'], resource['id'])
        provider, key, mime = resource['provider'], self.resource_id(resource['id']), resource['mime']
        if mime == 'application/vnd.google-apps.document':
            return {'kind': 'pdf', 'content': self.fetch(provider, '/files/' + key + '/export', {'mimeType': 'application/pdf'}, True)}
        if provider == 'microsoft' and resource['name'].lower().endswith('.docx'):
            try:
                converted = self.fetch(provider, '/me/drive/items/' + key + '/content', {'format': 'pdf'}, True)
                if converted.startswith(b'%PDF-'):
                    return {'kind': 'pdf', 'content': converted}
            except HTTPException as error:
                if error.status_code != 502:
                    raise
                # Conversion can be unavailable; retain the safe text preview.
        if mime == 'application/vnd.google-apps.spreadsheet':
            data = self.fetch(provider, '/files/' + key + '/export', {'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}, True)
        else:
            data = self.fetch(provider, '/files/' + key, {'alt': 'media'}, True) if provider == 'google' else self.fetch(provider, '/me/drive/items/' + key + '/content', binary=True)
        if resource['name'].lower().endswith('.pdf') or mime == 'application/pdf':
            return {'kind': 'pdf', 'content': data}
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024 or len(archive.infolist()) > 5000:
                    raise HTTPException(413, 'This document is too large to preview safely')
                if resource['kind'] == 'document':
                    xml = fromstring(archive.read('word/document.xml'))
                    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    paragraphs = [''.join(text.text or '' for text in paragraph.findall('.//w:t', namespace)) for paragraph in xml.findall('.//w:p', namespace)]
                    return {'kind': 'document', 'paragraphs': paragraphs[:2000], 'truncated': len(paragraphs) > 2000}
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                sheets = [{'name': sheet.title, 'rows': [[str(cell) if cell is not None else '' for cell in row] for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 200, 200), max_col=min(sheet.max_column or 30, 30), values_only=True)], 'truncated': (sheet.max_row or 0) > 200 or (sheet.max_column or 0) > 30} for sheet in workbook.worksheets[:20]]
            finally:
                workbook.close()
            return {'kind': 'sheet', 'sheets': sheets, 'truncated': len(workbook.worksheets) > 20}
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, 'This file could not be previewed. Try a standard .docx or .xlsx file.')
